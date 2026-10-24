import pymupdf
from docx import Document
from io import BytesIO



def extract_text_from_pdf(file):
    text = ""

    pdf_bytes = file.getvalue()

    with pymupdf.open(
    stream=pdf_bytes,
    filetype="pdf"
    ) as pdf:

        for page in pdf:
            text += page.get_text(sort=True) + "\n"

    return text.strip()


def extract_text_from_docx(file):
    text = []

    docx_bytes = BytesIO(file.getvalue())
    document = Document(docx_bytes)

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)

    return "\n".join(text).strip()


def extract_resume_text(file):

    if file is None:
        return ""

    filename = file.name.lower()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file)

    elif filename.endswith(".docx"):
        return extract_text_from_docx(file)

    else:
        raise ValueError("Unsupported file format")